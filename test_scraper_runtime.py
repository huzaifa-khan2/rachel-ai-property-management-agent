import os
import json
import tempfile
import unittest
from unittest.mock import patch

import openpyxl


for key in (
    "APPFOLIO_EMAIL",
    "APPFOLIO_PASSWORD",
    "RACHEL_GMAIL_USER",
    "RACHEL_GMAIL_PASSWORD",
):
    os.environ.setdefault(key, "deterministic-test-value")


import rachel_scraper_v2 as scraper


class _Element:
    def __init__(self, on_click=None):
        self.on_click = on_click

    def click(self):
        if self.on_click:
            self.on_click()

    def is_visible(self):
        return True


class _ExportPage:
    def __init__(self, create_export=None):
        self.actions = _Element()
        self.export = _Element(create_export)

    def goto(self, *_args, **_kwargs):
        return None

    def query_selector(self, selector):
        if "Actions" in selector:
            return self.actions
        if "Export as Excel" in selector:
            return self.export
        return None

    def query_selector_all(self, _selector):
        return []


class _RentRollPage:
    def __init__(self, ledger_snapshots):
        self.url = "https://example.test/dashboard"
        self.ledger_snapshots = list(ledger_snapshots)
        self.goto_calls = 0
        self.scroll_calls = 0

    def goto(self, url, **_kwargs):
        self.goto_calls += 1
        self.url = url

    def evaluate(self, _script, argument=None):
        if isinstance(argument, str):
            if self.ledger_snapshots:
                return self.ledger_snapshots.pop(0)
            return []
        self.scroll_calls += 1
        return None


class ScraperRuntimeTests(unittest.TestCase):
    def test_tenant_directory_parses_properties_units_and_residents(self):
        with tempfile.TemporaryDirectory() as root:
            workbook_path = os.path.join(root, "tenant-directory.xlsx")
            headers = [f"Synthetic Column {index}" for index in range(57)]
            for index, name in {
                0: "Property", 1: "Unit", 2: "Tenant", 3: "Status",
                4: "Tenant Type", 8: "Lease To", 9: "Rent",
                27: "Lease From", 30: "Eligible for Rent Increase",
                31: "Last Rent Increase", 32: "Next Rent Increase Date",
                50: "Insurance Provider", 51: "Insurance Expiration",
                52: "Insurance Policy Number", 54: "Primary Tenant",
            }.items():
                headers[index] = name

            def row(property_value, unit, tenant, primary, tenant_type):
                values = [""] * len(headers)
                values[0] = property_value
                values[1] = unit
                values[2] = tenant
                values[3] = "Current"
                values[4] = tenant_type
                values[8] = "08/31/2027"
                values[9] = "$1,250.00"
                values[27] = "09/01/2026"
                values[30] = "09/01/2027"
                values[50] = "Synthetic Insurance"
                values[51] = "12/31/2027"
                values[52] = "SYNTHETIC-POLICY"
                values[54] = "Yes" if primary else "No"
                return values

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(headers)
            sheet.append(row(
                "Synthetic Property - 100 Example Street", "1",
                "Resident, Primary", True, "Financially Responsible",
            ))
            sheet.append(row(
                "Synthetic Property - 100 Example Street", "1",
                "Resident, Occupant", False, "Other Occupant",
            ))
            sheet.append(row(
                "Synthetic Property - 100 Example Street", "2",
                "Resident, Second", True, "Financially Responsible",
            ))
            workbook.save(workbook_path)
            workbook.close()

            records = scraper.parse_tenant_directory_excel(workbook_path)

            self.assertEqual(len(records), 2)
            first = records[("100 Example Street", "1")]
            self.assertEqual(first["primary_tenant"], "Primary Resident")
            self.assertEqual(
                json.loads(first["all_tenants"]),
                ["Primary Resident", "Occupant Resident"],
            )
            self.assertEqual(first["current_rent"], 1250.0)
            self.assertEqual(first["lease_type"], "fixed")
            self.assertEqual(
                first["insurance"]["provider"], "Synthetic Insurance"
            )

    def test_rent_roll_adapter_accumulates_and_preserves_ambiguity(self):
        page = _RentRollPage([
            [
                {
                    "href": "?party_ids%5D=t_101",
                    "text": "Resident, Unique",
                },
                {
                    "href": "?party_ids%5D=t_201",
                    "text": "Resident, Duplicate",
                },
            ],
            [
                {
                    "href": "?party_ids%5D=t_202",
                    "text": "Resident, Duplicate",
                },
                {
                    "href": "?party_ids%5D=t_999",
                    "text": "Vacant-Unrented",
                },
            ],
        ])
        with patch.object(scraper, "RENT_ROLL_URL", (
            "https://example.test/buffered_reports/rent_roll"
        )), patch.object(scraper, "human_delay"), patch.object(
            scraper.time, "sleep"
        ):
            result = scraper.get_occupancy_ids(page)

        self.assertEqual(result["Resident, Unique"], "101")
        self.assertIsNone(result["Resident, Duplicate"])
        self.assertNotIn("Vacant-Unrented", result)
        self.assertEqual(page.goto_calls, 1)
        self.assertGreaterEqual(page.scroll_calls, 1)

    def test_export_cleanup_is_scoped_to_canonical_download_directory(self):
        with tempfile.TemporaryDirectory() as export_root, \
                tempfile.TemporaryDirectory() as outside_root:
            inside = os.path.join(export_root, "temporary.xlsx")
            outside = os.path.join(outside_root, "must-remain.xlsx")
            for path in (inside, outside):
                with open(path, "wb") as file_handle:
                    file_handle.write(b"temporary")

            with patch.object(
                scraper, "get_downloads_dir", return_value=export_root
            ):
                self.assertTrue(scraper.cleanup_temporary_export(
                    inside, "Test"
                ))
                self.assertFalse(scraper.cleanup_temporary_export(
                    outside, "Test"
                ))

            self.assertFalse(os.path.exists(inside))
            self.assertTrue(os.path.exists(outside))

    def test_invalid_tenant_export_fails_without_runtime_package_install(self):
        with tempfile.TemporaryDirectory() as root:
            workbook_path = os.path.join(root, "invalid.xlsx")
            with open(workbook_path, "wb") as file_handle:
                file_handle.write(b"not-an-excel-workbook")

            with patch.object(scraper.os, "system") as system_call:
                self.assertEqual(
                    scraper.parse_tenant_directory_excel(workbook_path),
                    {},
                )

            system_call.assert_not_called()
            self.assertFalse(os.path.exists(
                workbook_path.replace(".xlsx", "_renamed.xlsx")
            ))

    def test_occupancy_matching_fails_closed_on_duplicate_or_fuzzy_ambiguity(self):
        self.assertEqual(
            scraper.match_occupancy_id(
                "Jamie Tenant", "100 Oak Street",
                {"Tenant, Jamie": "101"},
            ),
            "101",
        )
        self.assertIsNone(scraper.match_occupancy_id(
            "Jamie Tenant", "100 Oak Street",
            {"Tenant, Jamie": None},
        ))
        self.assertIsNone(scraper.match_occupancy_id(
            "Unknown Tenant", "100 Oak Street",
            {"Tenant, Jamie": "101", "Tenant, Morgan": "102"},
        ))

    def test_export_uses_canonical_download_directory_and_new_file_only(self):
        with tempfile.TemporaryDirectory() as root:
            stale = os.path.join(root, "prior.xlsx")
            with open(stale, "wb") as file_handle:
                file_handle.write(b"stale")

            def create_export():
                with open(os.path.join(root, "playwrightdownload"), "wb") as handle:
                    handle.write(b"x" * 10001)

            with patch.object(
                scraper, "get_downloads_dir", return_value=root
            ), patch.object(scraper, "human_delay"), patch.object(
                scraper.time, "sleep"
            ):
                result = scraper.export_report_to_excel(
                    _ExportPage(create_export),
                    "https://example.test/buffered_reports/test-report",
                    "Test",
                )

            self.assertEqual(result, os.path.join(root, "playwrightdownload.xlsx"))
            self.assertTrue(os.path.exists(result))
            self.assertTrue(os.path.exists(stale))

    def test_timed_out_export_never_reuses_stale_report(self):
        with tempfile.TemporaryDirectory() as root:
            stale = os.path.join(root, "prior.xlsx")
            with open(stale, "wb") as file_handle:
                file_handle.write(b"x" * 10001)

            with patch.object(
                scraper, "get_downloads_dir", return_value=root
            ), patch.object(scraper, "human_delay"), patch.object(
                scraper.time, "sleep"
            ):
                result = scraper.export_report_to_excel(
                    _ExportPage(),
                    "https://example.test/buffered_reports/test-report",
                    "Test",
                )

            self.assertIsNone(result)

    def test_work_order_history_import_is_incremental_across_restart(self):
        with tempfile.TemporaryDirectory() as root:
            workbook_path = os.path.join(root, "work-orders.xlsx")
            database_path = os.path.join(root, "rachel.db")
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append([
                "Property", "Vendor", "Work Order Issue", "Status",
                "Work Order Number", "Job Description", "Work Done On",
                "Amount",
            ])
            sheet.append([
                "100 - 100 Oak Street", "Trusted Plumbing", "Plumbing",
                "Completed", "4000-1", "Repair fixture", "2026-08-01",
                125.0,
            ])
            workbook.save(workbook_path)
            workbook.close()

            self.assertTrue(scraper.parse_wo_report_and_update_vendors(
                workbook_path, database_path
            ))
            self.assertTrue(scraper.parse_wo_report_and_update_vendors(
                workbook_path, database_path
            ))

            import sqlite3
            conn = sqlite3.connect(database_path)
            try:
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM vendor_work_history "
                        "WHERE wo_num = '4000-1'"
                    ).fetchone()[0],
                    1,
                )
                self.assertEqual(
                    conn.execute(
                        "SELECT COUNT(*) FROM property_vendors"
                    ).fetchone()[0],
                    1,
                )
            finally:
                conn.close()

    def test_work_order_history_requires_identity_columns(self):
        with tempfile.TemporaryDirectory() as root:
            workbook_path = os.path.join(root, "invalid-work-orders.xlsx")
            database_path = os.path.join(root, "rachel.db")
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append([
                "Property", "Vendor", "Status", "Job Description",
                "Amount", "Extra One", "Extra Two",
            ])
            sheet.append([
                "100 Oak Street", "Trusted Plumbing", "Completed",
                "Repair fixture", 125.0, "", "",
            ])
            workbook.save(workbook_path)
            workbook.close()

            self.assertFalse(scraper.parse_wo_report_and_update_vendors(
                workbook_path, database_path
            ))


if __name__ == "__main__":
    unittest.main()
